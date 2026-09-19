"""Baixa as planilhas de dados abertos da EPE e mostra a estrutura de cada aba (para configurar ingestao/epe.py)."""
import sys
import time
from pathlib import Path

import curl_cffi.requests as cr
import openpyxl

BASE = Path(__file__).resolve().parent.parent
UA = {"User-Agent": "Mozilla/5.0 (compatible; FlexIA-Coletor/1.0; +hackathon ONS)"}
RAIZ = "https://www.epe.gov.br/sites-pt/publicacoes-dados-abertos/dados-abertos/Documents/"
ARQUIVOS = ["Dados_abertos_Consumo_Mensal.xlsx", "Dados_abertos_Mercado_Distribuicao.xlsx", "Dados%20brutos.xlsx",
            "PDE%202035_Painel%20de%20Resultados_Dados%20Abertos.xlsx"]

pasta = BASE / "out" / "fontes" / "epe"
pasta.mkdir(parents=True, exist_ok=True)
for a in ARQUIVOS:
    destino = pasta / a.replace("%20", "_")
    if not destino.exists():
        r = cr.get(RAIZ + a, timeout=120, impersonate="chrome", headers=UA)
        destino.write_bytes(r.content)
        time.sleep(2)
    wb = openpyxl.load_workbook(destino, read_only=True, data_only=True)
    print(f"\n#### {destino.name} ({destino.stat().st_size / 1e6:.1f} MB)")
    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(max_row=4, values_only=True))
        print(f"  [{ws.title}] ~{ws.max_row} linhas x {ws.max_column} colunas")
        for l in linhas:
            print("     ", [c for c in l][:12])
    sys.stdout.flush()
